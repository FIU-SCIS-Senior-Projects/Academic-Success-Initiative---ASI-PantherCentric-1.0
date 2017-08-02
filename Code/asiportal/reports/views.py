from django.contrib.auth.models import User
from django.utils import timezone
from django.http import HttpResponse
from django.shortcuts import render
from django.views import generic
from .services import (get_overall_weeks_stats,
        get_session_per_course_week_stats,
        get_course_week_stats,
        get_ambassador_stats,
        get_average_survey_ratings,
        get_missing_amb_surveys,
        get_missing_tut_surveys,
        get_survey_per_week_stats,
        add_data,
        add_charts,
        )
from semesters.models import Semester
from surveys.models import AmbassadorSurvey, TuteeSurvey
from tutoring_sessions.models import Session, IndividualSession
from collections import defaultdict
from openpyxl import load_workbook

class DirectoryView(generic.TemplateView):
    template_name = 'reports/directory.html'

class StatisticPrepare(generic.TemplateView):
    stats_function = None

    def __init__(self, *args, **kwargs):
        super(StatisticPrepare, self).__init__(*args, **kwargs)
        current_semester = Semester.objects.current()
        if current_semester.count() > 1:
            self.sem = current_semester.filter(term='SC')[0]
        else:
            self.sem = current_semester[0]

    def get_stats_function(self):
        return self.stats_function

    def get_context_data(self, *args, **kwargs):
        cd = super(StatisticPrepare, self).get_context_data(*args, **kwargs)
        cd['cs'] = self.sem
        cd['weeks'] = self.get_stats_function()(self.sem)
        return cd


class WeekToWeekView(StatisticPrepare):
    template_name = 'reports/week_to_week.html'
    stats_function = get_overall_weeks_stats

class WeekToWeekCourseView(StatisticPrepare):
    template_name = 'reports/week_to_week_course.html'
    stats_function = get_course_week_stats

class WeekToWeekSessionView(StatisticPrepare):
    template_name = 'reports/week_to_week_session.html'
    stats_function = get_session_per_course_week_stats

class WeekToWeekSurveyView(StatisticPrepare):
    template_name = 'reports/week_to_week_survey.html'
    stats_function = get_survey_per_week_stats

class AmbassadorAverageRatingsView(StatisticPrepare):
    template_name = 'reports/average_ratings.html'
    stats_function = get_average_survey_ratings

class AmbassadorReportsView(generic.ListView):
    model = User
    template_name = 'reports/ambassador_list.html'

    def get_queryset(self):
        qs = User.objects.filter(groups__name='Ambassador')
        return qs

class IndividualAmbassadorReportView(generic.DetailView):
    model = User
    template_name = 'reports/ind_ambassador_stats.html'

    def __init__(self, *args, **kwargs):
        super(IndividualAmbassadorReportView, self).__init__(*args, **kwargs)
        current_semester = Semester.objects.current()
        if current_semester.count() > 1:
            self.sem = current_semester.filter(term='SC')[0]
        else:
            self.sem = current_semester[0]

    def get_context_data(self, *args, **kwargs):
        cd = super(IndividualAmbassadorReportView, self).get_context_data(*args, **kwargs)
        cd['weeks'] = get_ambassador_stats(self.object, self.sem)
        return cd

class TestGraphsView(generic.TemplateView):
    template_name = 'reports/test_graph.html'

class MissingSurveyReportView(generic.TemplateView):
    template_name = 'reports/missing_survey_report.html'
    def get_context_data(self, *args, **kwargs):
        cd = super(MissingSurveyReportView, self).get_context_data(*args, **kwargs)
        cd['tuts_survs'] = get_missing_tut_surveys()
        cd['amb_survs'] = get_missing_amb_surveys()
        return cd



'''
SPRINT 3 SECTION:
    Eerything goes below this comment.
'''
class SemesterReportView(generic.TemplateView):
    template_name = 'reports/semester_report.html'
    
    def get_context_data(self, *args, **kwargs):
        cd = super(SemesterReportView, self).get_context_data(*args, **kwargs)
		
		#Sessions Table information
        scheduledSessions = IndividualSession.objects.all().count()
        actualSessions = AmbassadorSurvey.objects.filter(tutee_absent= False).count()
        cd['session_info'] = {"actual": actualSessions, "scheduled":scheduledSessions}
        
        #get course info hard coded
        cd['course_info'] = [
            {"course_name": "COP4338", "number_of_sessions": "15", "percentage_of_total" : "25"},
            {"course_name": "COP4338", "number_of_sessions": "15", "percentage_of_total" : "25"},
            {"course_name": "COP4338", "number_of_sessions": "15", "percentage_of_total" : "25"},
            {"course_name": "COP4338", "number_of_sessions": "15", "percentage_of_total" : "25"},
        ]
		
        #Tutee Table Information
        sessions = Session.objects.prefetch_related('tutee', 'course').all()
        courseList = defaultdict(list)
        for session in sessions:
            courseList[session.tutee] +=  [session.course.name] if session.course not in courseList[session.tutee] else []
        cd['tutee_info'] = []

        for i in courseList.keys():
            obj = User.objects.get(username=i)
            tuteeSessions = Session.objects.filter(tutee=i)
            tuteeSingleSessions = []
            tuteeAbsences = 0
            tuteeSurveys = 0
            for tuteeSession in tuteeSessions:
                tuteeSingleSessions += IndividualSession.objects.filter(session=tuteeSession)
            for singleSession in tuteeSingleSessions:
                tuteeAbsences += AmbassadorSurvey.objects.filter(individual_session=singleSession).filter(tutee_absent=True).count()
                tuteeSurveys += TuteeSurvey.objects.filter(individual_session=singleSession).filter(submitted=False).count()

            cd['tutee_info'].append({"name": obj.get_full_name,
                                "scheduled": len(tuteeSingleSessions),
                                "actual": len(tuteeSingleSessions) - tuteeAbsences,
                                "courses": ", ".join(courseList.get(i)),
                                "survey": (len(tuteeSingleSessions) - tuteeSurveys) / len(tuteeSingleSessions) * 100 if len(tuteeSingleSessions) > 0 else 0}
                                )

        #get the other things below or above
        
        #Ambassador Table
        # cd['Ambassador_info'] = []
        # obj = User.objects.get(username=i)
        
        
        return cd

def download_report_excel(request):
    wb = load_workbook('reports/templates/reports/report_template.xlsx')
    response = HttpResponse(content_type='application/ms-excel')
    # http pls
    response['Content-Disposition'] = 'attachment; filename=report-{}.xlsx'.format(timezone.now())
    wb = add_data(wb)
    wb = add_charts(wb)
    wb._sheets = [wb._sheets[1], wb._sheets[0]]
    wb.save(response)
    return response
